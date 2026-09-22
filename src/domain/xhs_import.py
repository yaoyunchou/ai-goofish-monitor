"""小红书商品 Excel 模板与行解析。说明放在第二张表，避免示例被导入。"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook, load_workbook

from src.domain.xhs_labels import split_tag_cell

IMPORT_HEADERS = ("商品链接", "店铺", "分类", "标记")
MAX_IMPORT_ROWS = 500


def build_import_template() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "商品"
    sheet.append(list(IMPORT_HEADERS))
    sheet.column_dimensions["A"].width = 48
    sheet.column_dimensions["B"].width = 24
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["D"].width = 28
    notes = book.create_sheet("说明")
    notes["A1"] = "从「商品」表第 2 行开始填写。空行会跳过。"
    notes["A2"] = "商品链接必填，可以是公开商品链接或商品 ID。"
    notes["A3"] = "店铺、分类可以空着。店铺没有时会按名字新建。"
    notes["A4"] = "标记写在一个格子里，用逗号或顿号分开，最多 8 个。"
    notes["A5"] = "已经在监控里的商品，空着的格子不会清掉原来的店铺、分类和标记。"
    notes.column_dimensions["A"].width = 72
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def parse_import_workbook(payload: bytes) -> list[dict]:
    if not payload:
        raise ValueError("表格是空的")
    try:
        book = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("无法读取 xlsx，请用下载的模板") from exc
    try:
        sheet = book["商品"] if "商品" in book.sheetnames else book.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        book.close()
    if not rows:
        raise ValueError("表格没有表头")
    header = tuple(_cell(value) for value in rows[0][: len(IMPORT_HEADERS)])
    if header != IMPORT_HEADERS:
        raise ValueError("表头必须是：商品链接、店铺、分类、标记")
    parsed: list[dict] = []
    for offset, raw in enumerate(rows[1:], start=2):
        url, shop, category, tags = (_cell(raw[index] if index < len(raw) else None) for index in range(4))
        if not any((url, shop, category, tags)):
            continue
        parsed.append(
            {
                "row": offset,
                "url": url,
                "shop": shop or None,
                "category": category or None,
                "tags": split_tag_cell(tags),
            }
        )
        if len(parsed) > MAX_IMPORT_ROWS:
            raise ValueError(f"单次最多 {MAX_IMPORT_ROWS} 行")
    return parsed


def _cell(value) -> str:
    if value is None:
        return ""
    return str(value).strip()
