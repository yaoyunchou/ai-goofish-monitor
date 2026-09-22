from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from io import BytesIO

from src.domain.xhs_import import IMPORT_HEADERS, build_import_template, parse_import_workbook
from src.domain.xhs_labels import normalize_tags, split_tag_cell
from src.domain.xhs_shops import rollup_products
from src.services.xhs_storage import (
    add_product,
    apply_collect_results,
    collect_products,
    get_product,
    ignore_failures,
    import_product_rows,
    list_board,
    list_products_by_status,
)
from src.xhs_collector import CollectOutcome, FetchResult


GOODS = "https://www.xiaohongshu.com/goods-detail/abcdef0123456789"
OTHER = "https://www.xiaohongshu.com/goods-detail/abcdef0123456780"


def test_tags_split_and_reject_too_many():
    assert split_tag_cell(" 热卖，新品、现货 ") == ["热卖", "新品", "现货"]
    assert split_tag_cell("   ") is None
    try:
        normalize_tags([f"标记{index}" for index in range(9)])
    except ValueError as exc:
        assert "8" in str(exc)
    else:
        raise AssertionError("expected too many tags")


def test_rollup_skips_missing_baseline_and_marks_incomplete():
    summary = rollup_products(
        [
            {"sold_total": 10, "today": None, "today_incomplete": True, "yesterday": 2, "yesterday_incomplete": False, "last_hour": None, "last_hour_incomplete": False},
            {"sold_total": None, "today": 3, "today_incomplete": True, "yesterday": None, "yesterday_incomplete": True, "last_hour": 1, "last_hour_incomplete": False},
        ]
    )
    assert summary["product_count"] == 2
    assert summary["sold_total"] == 10
    assert summary["today"] == 3
    assert summary["today_incomplete"] is True
    assert summary["yesterday"] == 2
    assert summary["yesterday_incomplete"] is False
    assert summary["last_hour"] == 1


def test_template_headers_and_blank_rows_are_skipped():
    payload = build_import_template()
    book = load_workbook(BytesIO(payload))
    assert tuple(cell.value for cell in book["商品"][1]) == IMPORT_HEADERS
    book.close()

    edited = Workbook()
    sheet = edited.active
    sheet.title = "商品"
    sheet.append(list(IMPORT_HEADERS))
    sheet.append(["", "", "", ""])
    sheet.append(["not-a-link", "", "", ""])
    buffer = BytesIO()
    edited.save(buffer)
    rows = parse_import_workbook(buffer.getvalue())
    assert rows == [{"row": 3, "url": "not-a-link", "shop": None, "category": None, "tags": None}]


def test_manual_shop_survives_page_name(offline_db):
    saved = add_product(GOODS, shop_name="手选店", category="家居", tags=["热卖"])
    apply_collect_results(
        [
            CollectOutcome(
                product_id=saved["id"],
                ok=True,
                sold=12,
                shop_name="页面店",
                captured_at=datetime.now(ZoneInfo("Asia/Shanghai")),
            )
        ]
    )
    product = get_product(saved["id"])
    assert product["assigned_shop_name"] == "手选店"
    assert product["shop_name"] == "页面店"
    assert product["category"] == "家居"
    assert product["tags"] == ["热卖"]


def test_same_shop_name_is_reused_and_blank_cells_do_not_clear(offline_db):
    first = add_product(GOODS, shop_name="同名店", category="家居", tags=["热卖"])
    second = add_product(OTHER, shop_name="同名店")
    assert first["shop_id"] == second["shop_id"]
    again = add_product(GOODS)
    assert again["created"] is False
    assert again["assigned_shop_name"] == "同名店"
    assert again["category"] == "家居"
    assert again["tags"] == ["热卖"]


def test_import_records_bad_links_and_keeps_existing_labels(offline_db):
    add_product(GOODS, shop_name="手选店", category="家居", tags=["热卖"])
    book = Workbook()
    sheet = book.active
    sheet.title = "商品"
    sheet.append(list(IMPORT_HEADERS))
    sheet.append([GOODS, "", "", ""])
    sheet.append(["bad-link", "新店", "", ""])
    buffer = BytesIO()
    book.save(buffer)
    result = import_product_rows(buffer.getvalue())
    assert result["added"] == 0
    assert result["updated"] == 1
    assert result["failed"] == [{"row": 3, "reason": "无法从链接解析商品 ID，请粘贴商品页链接或商品 ID"}]
    product = get_product("abcdef0123456789")
    assert product["assigned_shop_name"] == "手选店"
    assert product["tags"] == ["热卖"]


def test_failures_and_delisted_leave_the_board(offline_db):
    saved = add_product(GOODS)
    other = add_product(OTHER)
    now = datetime.now(ZoneInfo("Asia/Shanghai"))
    apply_collect_results(
        [
            CollectOutcome(product_id=saved["id"], ok=False, error="超时", captured_at=now),
            CollectOutcome(product_id=other["id"], ok=False, delisted=True, error="商品已下架", captured_at=now),
        ]
    )
    assert list_board() == []
    assert [row["id"] for row in list_products_by_status(("failed", "skipped"))] == [saved["id"]]
    assert list_products_by_status(("delisted",))[0]["last_error"] == "商品已下架"
    assert ignore_failures([saved["id"]]) == 1
    assert [row["id"] for row in list_board()] == [saved["id"]]
    calls = []

    def fetch(url: str) -> FetchResult:
        calls.append(url)
        return FetchResult(status=200, final_url=url, body='"soldCount": 3')

    collect_products(fetch=fetch)
    assert calls == [GOODS]
